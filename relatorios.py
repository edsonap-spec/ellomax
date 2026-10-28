"""Exportação para Excel (openpyxl) e PDF (reportlab)."""
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

FMT_BRL = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
FMT_NUM = '#,##0.##'
FMT_PCT = '0.0"%"'


def _dt(v):
    try:
        return datetime.strptime(v[:10], "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return v or None


def _brl(v):
    s = f"{abs(v or 0):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("-R$ " if (v or 0) < -0.004 else "R$ ") + s


def _num(v):
    v = v or 0
    if float(v).is_integer():
        return f"{int(v):,}".replace(",", ".")
    return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _d(v):
    d = _dt(v)
    return d.strftime("%d/%m/%Y") if isinstance(d, date) else (d or "")


# ---------------------------------------------------------------- Excel

def _aba(wb, titulo, cabecalho, linhas, formatos=None, primeira=False):
    ws = wb.active if primeira else wb.create_sheet()
    ws.title = titulo[:31]
    ws.append(cabecalho)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0D6EFD")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for li in linhas:
        ws.append(list(li))
    formatos = formatos or {}
    for idx, nome in enumerate(cabecalho, start=1):
        letra = get_column_letter(idx)
        fmt = formatos.get(nome)
        largura = max([len(str(nome))] + [len(str(li[idx - 1])) for li in linhas if li[idx - 1] is not None] or [10])
        ws.column_dimensions[letra].width = min(max(largura + 2, 10), 60)
        for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            for c in row:
                if fmt:
                    c.number_format = fmt
                elif isinstance(c.value, date):
                    c.number_format = "DD/MM/YYYY"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def _salvar(wb):
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def xlsx_editais(linhas):
    wb = Workbook()
    cab = ["Empresa", "Edital", "Processo", "Contrato", "Órgão", "Modalidade", "Objeto", "Situação",
           "Vigência início", "Vigência fim", "Contratado", "Empenhado", "Entregue", "Recebido + retido",
           "A receber", "Em atraso", "Gastos", "Impostos estimados", "Lucro projetado", "Margem %"]
    dados = [(e["empresa"], e["numero"], e["processo"], e["numero_contrato"], e["orgao"], e["modalidade"], e["objeto"],
              e["status"], _dt(e["vigencia_inicio"]), _dt(e["vigencia_fim"]), r["contratado"], r["empenhado"], r["entregue"],
              r["liquidado"], r["a_receber"], r["atrasado"], r["gastos"], r["impostos_estimados"], r["lucro_projetado"],
              r["margem_projetada"]) for e, r in linhas]
    moeda = {k: FMT_BRL for k in cab[10:19]}
    moeda["Margem %"] = FMT_PCT
    _aba(wb, "Editais", cab, dados, moeda, primeira=True)
    return _salvar(wb)


def xlsx_edital(d):
    e, r = d["e"], d["r"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    campos = [("Empresa", e["empresa"]), ("Órgão", e["orgao"]), ("CNPJ do órgão", e["cnpj_orgao"]), ("UASG", e["uasg"]),
              ("Nº do edital", e["numero"]), ("Processo", e["processo"]), ("Modalidade", e["modalidade"]),
              ("Contrato / ata", e["numero_contrato"]), ("Situação", e["status"]), ("Objeto", e["objeto"]),
              ("Abertura", _dt(e["data_abertura"])), ("Homologação", _dt(e["data_homologacao"])),
              ("Vigência início", _dt(e["vigencia_inicio"])), ("Vigência fim", _dt(e["vigencia_fim"])),
              ("Prazo de entrega", e["prazo_entrega"]), ("Local de entrega", e["local_entrega"]),
              ("Contato no órgão", e["contato_orgao"]), ("Responsável", e["responsavel"]), (None, None),
              ("Valor contratado", r["contratado"]), ("  dos itens", r["contratado_itens"]), ("  aditivos", r["aditivos"]),
              ("Empenhado", r["empenhado"]), ("Entregue", r["entregue"]), ("Recebido (líquido)", r["recebido"]),
              ("Retenções", r["retencoes"]), ("Saldo a receber", r["a_receber"]), ("Em atraso", r["atrasado"]),
              ("Gastos lançados", r["gastos"]), ("Custo estimado dos itens", r["custo_estimado"]),
              ("Impostos estimados", r["impostos_estimados"]), ("Lucro projetado", r["lucro_projetado"])]
    for rot, val in campos:
        ws.append([rot, val])
        if isinstance(val, float):
            ws.cell(ws.max_row, 2).number_format = FMT_BRL
        elif isinstance(val, date):
            ws.cell(ws.max_row, 2).number_format = "DD/MM/YYYY"
    ws.append(["Margem projetada", r["margem_projetada"]])
    ws.cell(ws.max_row, 2).number_format = FMT_PCT
    for c in ws["A"]:
        c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 80

    _aba(wb, "Itens", ["Lote", "Item", "Descrição", "Marca", "Unidade", "Quantidade", "Valor unitário", "Total",
                       "Empenhado", "Entregue", "Saldo", "Custo unitário"],
         [(i["lote"], i["numero_item"], i["descricao"], i["marca"], i["unidade"], i["quantidade"], i["valor_unitario"],
           i["quantidade"] * i["valor_unitario"], i["empenhado"], i["qtd_entregue"], i["quantidade"] - i["qtd_entregue"],
           i["custo_unitario"]) for i in d["itens"]],
         {"Valor unitário": FMT_BRL, "Total": FMT_BRL, "Custo unitário": FMT_BRL, "Quantidade": FMT_NUM,
          "Empenhado": FMT_NUM, "Entregue": FMT_NUM, "Saldo": FMT_NUM})
    _aba(wb, "Empenhos e OF", ["Tipo", "Número", "Emissão", "Prazo de entrega", "Entregue em", "Valor", "Itens", "Observações"],
         [(m["tipo"], m["numero"], _dt(m["data_emissao"]), _dt(m["prazo_entrega"]), _dt(m["data_entrega"]),
           m["valor_itens"] or m["valor"],
           "; ".join(f"item {li['numero_item'] or '-'}: {_num(li['quantidade'])}" for li in d["linhas_emp"].get(m["id"], [])),
           m["observacoes"]) for m in d["empenhos"]], {"Valor": FMT_BRL})
    _aba(wb, "Recebimentos", ["Descrição", "Empenho/OF", "NF", "Previsto em", "Valor previsto", "Recebido em",
                              "Valor recebido", "Retenções", "Diferença", "Observações"],
         [(x["descricao"], x["empenho"], x["nota_fiscal"], _dt(x["data_prevista"]), x["valor_previsto"],
           _dt(x["data_recebimento"]), x["valor_recebido"], x["retencoes"],
           x["valor_previsto"] - x["valor_recebido"] - x["retencoes"], x["observacoes"]) for x in d["receb"]],
         {"Valor previsto": FMT_BRL, "Valor recebido": FMT_BRL, "Retenções": FMT_BRL, "Diferença": FMT_BRL})
    _aba(wb, "Gastos", ["Data", "Categoria", "Descrição", "Fornecedor", "Documento", "Valor"],
         [(_dt(x["data"]), x["categoria"], x["descricao"], x["fornecedor"], x["documento"], x["valor"]) for x in d["gastos"]],
         {"Valor": FMT_BRL})
    _aba(wb, "Aditivos", ["Número", "Tipo", "Assinatura", "Descrição", "Nova vigência", "Valor", "Percentual"],
         [(x["numero"], x["tipo"], _dt(x["data_assinatura"]), x["descricao"], _dt(x["nova_vigencia_fim"]), x["valor"],
           x["percentual"]) for x in d["aditivos"]], {"Valor": FMT_BRL, "Percentual": FMT_PCT})
    _aba(wb, "Garantias", ["Modalidade", "Descrição", "%", "Valor", "Apólice", "Seguradora/banco", "Início", "Fim",
                           "Devolvida em"],
         [(x["tipo"], x["descricao"], x["percentual"], x["valor"], x["apolice"], x["seguradora"], _dt(x["data_inicio"]),
           _dt(x["data_fim"]), _dt(x["data_devolucao"])) for x in d["garantias"]], {"Valor": FMT_BRL, "%": FMT_PCT})
    return _salvar(wb)


# ---------------------------------------------------------------- PDF

def pdf_edital(d):
    e, r = d["e"], d["r"]
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm,
                            title=f"Edital {e['numero']}", author=e["empresa"] or "")
    est = getSampleStyleSheet()
    peq = est["BodyText"].clone("peq", fontSize=8, leading=10)
    h2 = est["Heading2"].clone("h2", spaceBefore=10, spaceAfter=4)

    def esc(t):
        return str(t if t is not None else "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")

    def P(marcacao):
        return Paragraph(marcacao, peq)

    def tabela(cab, linhas, larguras=None, total=None, direita=()):
        dados = [[P(f"<b>{esc(c)}</b>") for c in cab]] + [[P(esc(c)) for c in li] for li in linhas]
        if total:
            dados.append([P(f"<b>{esc(c)}</b>") for c in total])
        t = Table(dados, colWidths=larguras, repeatRows=1)
        estilo = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E7F0FF")),
                  ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BBBBBB")),
                  ("VALIGN", (0, 0), (-1, -1), "TOP")]
        for col in direita:
            estilo.append(("ALIGN", (col, 0), (col, -1), "RIGHT"))
        if total:
            estilo.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F2F2F2")))
        t.setStyle(TableStyle(estilo))
        return t

    hist = [Paragraph(esc(f"Edital {e['numero']} — {e['orgao']}"), est["Title"]),
            P(esc(f"{e['empresa'] or ''} · Situação: {e['status']} · Emitido em {date.today():%d/%m/%Y}")),
            Spacer(1, 6)]
    info = [("Modalidade", e["modalidade"]), ("Processo", e["processo"]), ("Contrato / ata", e["numero_contrato"]),
            ("CNPJ do órgão", e["cnpj_orgao"]), ("UASG", e["uasg"]), ("Homologação", _d(e["data_homologacao"])),
            ("Vigência", f"{_d(e['vigencia_inicio'])} a {_d(e['vigencia_fim'])}"), ("Prazo de entrega", e["prazo_entrega"]),
            ("Local de entrega", e["local_entrega"]), ("Contato", e["contato_orgao"]), ("Responsável", e["responsavel"]),
            ("Abertura", _d(e["data_abertura"]))]
    linhas_info = []
    for i in range(0, len(info), 3):
        li = []
        for rot, val in info[i:i + 3]:
            li += [f"<b>{esc(rot)}</b>", esc(val or "—")]
        linhas_info.append(li)
    t = Table([[P(c) for c in li] for li in linhas_info], colWidths=[28 * mm, 62 * mm] * 3)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#DDDDDD"))]))
    hist += [t, Spacer(1, 4), P(f"<b>Objeto:</b> {esc(e['objeto'])}")]

    hist.append(Paragraph("Resumo financeiro", h2))
    resumo = [("Contratado", r["contratado"]), ("Empenhado", r["empenhado"]), ("Entregue", r["entregue"]),
              ("Recebido + retido", r["liquidado"]), ("Saldo a receber", r["a_receber"]), ("Em atraso", r["atrasado"]),
              ("Gastos", r["gastos"]), ("Impostos estimados", r["impostos_estimados"]),
              ("Lucro projetado", r["lucro_projetado"])]
    hist.append(tabela([x[0] for x in resumo] + ["Margem"], [[_brl(x[1]) for x in resumo] + [f"{r['margem_projetada']:.1f}%"]],
                       direita=range(10)))

    if d["itens"]:
        hist.append(Paragraph("Itens vencidos", h2))
        hist.append(tabela(["Lote", "Item", "Descrição", "Marca", "Un.", "Qtd.", "Vlr. unit.", "Total", "Empenhado", "Entregue", "Saldo"],
                           [(i["lote"], i["numero_item"], i["descricao"], i["marca"], i["unidade"], _num(i["quantidade"]),
                             _brl(i["valor_unitario"]), _brl(i["quantidade"] * i["valor_unitario"]), _num(i["empenhado"]),
                             _num(i["qtd_entregue"]), _num(i["quantidade"] - i["qtd_entregue"])) for i in d["itens"]],
                           [12 * mm, 12 * mm, 80 * mm, 30 * mm, 12 * mm, 18 * mm, 24 * mm, 26 * mm, 20 * mm, 20 * mm, 18 * mm],
                           total=["", "", "Total", "", "", "", "", _brl(r["contratado_itens"]), "", "", ""],
                           direita=range(5, 11)))
    if d["empenhos"]:
        hist.append(Paragraph("Empenhos e ordens de fornecimento", h2))
        hist.append(tabela(["Tipo", "Número", "Emissão", "Prazo", "Entregue em", "Valor"],
                           [(m["tipo"], m["numero"], _d(m["data_emissao"]), _d(m["prazo_entrega"]), _d(m["data_entrega"]),
                             _brl(m["valor_itens"] or m["valor"])) for m in d["empenhos"]], direita=(5,)))
    if d["receb"]:
        hist.append(Paragraph("Recebimentos", h2))
        hist.append(tabela(["Descrição", "Empenho/OF", "NF", "Previsto em", "Valor previsto", "Recebido em", "Valor recebido",
                            "Retenções", "Diferença"],
                           [(x["descricao"], x["empenho"], x["nota_fiscal"], _d(x["data_prevista"]), _brl(x["valor_previsto"]),
                             _d(x["data_recebimento"]), _brl(x["valor_recebido"]), _brl(x["retencoes"]),
                             _brl(x["valor_previsto"] - x["valor_recebido"] - x["retencoes"])) for x in d["receb"]],
                           total=["Total", "", "", "", _brl(r["previsto"]), "", _brl(r["recebido"]), _brl(r["retencoes"]),
                                  _brl(r["previsto"] - r["liquidado"])],
                           direita=(4, 6, 7, 8)))
    if d["gastos"]:
        hist.append(Paragraph("Gastos", h2))
        hist.append(tabela(["Data", "Categoria", "Descrição", "Fornecedor", "Documento", "Valor"],
                           [(_d(x["data"]), x["categoria"], x["descricao"], x["fornecedor"], x["documento"], _brl(x["valor"]))
                            for x in d["gastos"]], total=["Total", "", "", "", "", _brl(r["gastos"])], direita=(5,)))
    if d["aditivos"]:
        hist.append(Paragraph("Termos aditivos", h2))
        hist.append(tabela(["Nº", "Tipo", "Assinatura", "Descrição", "Nova vigência", "Valor", "%"],
                           [(x["numero"], x["tipo"], _d(x["data_assinatura"]), x["descricao"], _d(x["nova_vigencia_fim"]),
                             _brl(x["valor"]), _num(x["percentual"]) if x["percentual"] is not None else "")
                            for x in d["aditivos"]], direita=(5, 6)))
    if d["garantias"]:
        hist.append(Paragraph("Garantias contratuais", h2))
        hist.append(tabela(["Modalidade", "Descrição", "Valor", "Apólice", "Seguradora/banco", "Vigência", "Devolvida em"],
                           [(x["tipo"], x["descricao"], _brl(x["valor"]), x["apolice"], x["seguradora"],
                             f"{_d(x['data_inicio'])} a {_d(x['data_fim'])}", _d(x["data_devolucao"])) for x in d["garantias"]],
                           direita=(2,)))
    doc.build(hist)
    buf.seek(0)
    return buf
